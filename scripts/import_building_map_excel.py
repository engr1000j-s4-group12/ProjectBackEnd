#!/usr/bin/env python3
from __future__ import annotations

import argparse
import csv
import json
import re
import sys
import zipfile
from pathlib import Path
from typing import Any
from xml.etree import ElementTree as ET


ROOT = Path(__file__).resolve().parents[1]
DEFAULT_INPUT = ROOT / "data" / "building_map_nodes.xlsx"
DEFAULT_OUTPUT = ROOT / "data" / "building_map.json"

DEFAULT_FLOOR_PLAN = {
    "floor": 4,
    "source": "龙宾楼4楼楼层索引图",
    "outline_dimensions_m": {
        "width_m": 70,
        "height_m": 80,
        "width_range_m": [65, 75],
        "height_range_m": [75, 85],
        "note": "长边约80m，短边约70m；坐标为按图片位置估算。",
    },
    "coordinate_system": {
        "origin": "floor_plan_top_left",
        "x_axis": "left_to_right",
        "y_axis": "top_to_bottom",
        "unit": "meter",
    },
}


def _split_items(value: Any) -> list[str]:
    if value is None:
        return []
    text = str(value).strip()
    if not text:
        return []
    parts = re.split(r"[|;；,，\n]+", text)
    return [part.strip() for part in parts if part.strip()]


def _to_int(value: Any, *, default: int | None = None) -> int | None:
    if value is None or str(value).strip() == "":
        return default
    try:
        return int(float(str(value).strip()))
    except ValueError as exc:
        raise ValueError(f"无法解析整数：{value!r}") from exc


def _to_float(value: Any, *, default: float | None = None) -> float | None:
    if value is None or str(value).strip() == "":
        return default
    try:
        return round(float(str(value).strip()), 4)
    except ValueError as exc:
        raise ValueError(f"无法解析数字：{value!r}") from exc


def _to_bool(value: Any, *, default: bool = True) -> bool:
    if value is None or str(value).strip() == "":
        return default
    text = str(value).strip().casefold()
    if text in {"true", "1", "yes", "y", "是", "可", "可以"}:
        return True
    if text in {"false", "0", "no", "n", "否", "不可", "不可以"}:
        return False
    raise ValueError(f"无法解析布尔值：{value!r}")


def _json_value(value: Any, *, default: Any) -> Any:
    if value is None or str(value).strip() == "":
        return default
    try:
        return json.loads(str(value).strip())
    except json.JSONDecodeError as exc:
        raise ValueError(f"无法解析 JSON：{value!r}") from exc


def _column_index(cell_ref: str) -> int:
    letters = re.sub(r"[^A-Z]", "", cell_ref.upper())
    index = 0
    for char in letters:
        index = index * 26 + ord(char) - ord("A") + 1
    return index - 1


def _xlsx_shared_strings(archive: zipfile.ZipFile) -> list[str]:
    try:
        payload = archive.read("xl/sharedStrings.xml")
    except KeyError:
        return []
    root = ET.fromstring(payload)
    namespace = {"x": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}
    strings: list[str] = []
    for item in root.findall("x:si", namespace):
        text_parts = [node.text or "" for node in item.findall(".//x:t", namespace)]
        strings.append("".join(text_parts))
    return strings


def _xlsx_sheet_path(archive: zipfile.ZipFile, sheet_name: str) -> str | None:
    workbook = ET.fromstring(archive.read("xl/workbook.xml"))
    namespace = {
        "x": "http://schemas.openxmlformats.org/spreadsheetml/2006/main",
        "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships",
    }
    relationship_id = None
    for sheet in workbook.findall(".//x:sheet", namespace):
        if sheet.attrib.get("name") == sheet_name:
            relationship_id = sheet.attrib[
                "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id"
            ]
            break
    if relationship_id is None:
        return None

    relationships = ET.fromstring(archive.read("xl/_rels/workbook.xml.rels"))
    rel_namespace = {
        "rel": "http://schemas.openxmlformats.org/package/2006/relationships"
    }
    for relationship in relationships.findall("rel:Relationship", rel_namespace):
        if relationship.attrib.get("Id") == relationship_id:
            target = relationship.attrib["Target"].lstrip("/")
            return target if target.startswith("xl/") else f"xl/{target}"
    return None


def _read_xlsx(path: Path, sheet_name: str) -> list[dict[str, str]]:
    try:
        import openpyxl  # type: ignore[import-not-found]
    except ModuleNotFoundError:
        openpyxl = None

    if openpyxl is not None:
        workbook = openpyxl.load_workbook(path, data_only=True)
        if sheet_name not in workbook.sheetnames:
            return []
        worksheet = workbook[sheet_name]
        rows = list(worksheet.iter_rows(values_only=True))
        if not rows:
            return []
        headers = [str(value).strip() if value is not None else "" for value in rows[0]]
        records: list[dict[str, str]] = []
        for row in rows[1:]:
            record = {
                header: "" if value is None else str(value).strip()
                for header, value in zip(headers, row)
                if header
            }
            if any(record.values()):
                records.append(record)
        return records

    with zipfile.ZipFile(path) as archive:
        shared_strings = _xlsx_shared_strings(archive)
        sheet_path = _xlsx_sheet_path(archive, sheet_name)
        if sheet_path is None:
            return []
        sheet_payload = archive.read(sheet_path)
    root = ET.fromstring(sheet_payload)
    namespace = {"x": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}
    matrix: list[list[str]] = []
    for row in root.findall(".//x:sheetData/x:row", namespace):
        values: list[str] = []
        for cell in row.findall("x:c", namespace):
            index = _column_index(cell.attrib.get("r", "A1"))
            while len(values) <= index:
                values.append("")
            value_node = cell.find("x:v", namespace)
            if value_node is None:
                inline_node = cell.find("x:is/x:t", namespace)
                values[index] = inline_node.text if inline_node is not None and inline_node.text else ""
                continue
            value = value_node.text or ""
            if cell.attrib.get("t") == "s":
                values[index] = shared_strings[int(value)] if value else ""
            else:
                values[index] = value
        matrix.append(values)
    if not matrix:
        return []
    headers = [value.strip() for value in matrix[0]]
    records = []
    for row in matrix[1:]:
        record = {
            header: row[index].strip() if index < len(row) else ""
            for index, header in enumerate(headers)
            if header
        }
        if any(record.values()):
            records.append(record)
    return records


def _read_csv(path: Path) -> list[dict[str, str]]:
    with path.open(encoding="utf-8-sig", newline="") as file:
        return [row for row in csv.DictReader(file) if any((value or "").strip() for value in row.values())]


def read_rows(path: Path, sheet_name: str) -> list[dict[str, str]]:
    suffix = path.suffix.lower()
    if suffix == ".xlsx":
        return _read_xlsx(path, sheet_name)
    if suffix == ".csv":
        return _read_csv(path)
    raise ValueError("仅支持 .xlsx 或 .csv 文件")


def _landmarks_from_row(row: dict[str, str], room_number: str) -> list[dict[str, Any]]:
    parsed = _json_value(row.get("visual_landmarks_json"), default=None)
    if parsed is not None:
        if not isinstance(parsed, list):
            raise ValueError("visual_landmarks_json 必须是数组")
        return parsed

    landmark_aliases = _split_items(row.get("visual_landmark_aliases"))
    if not landmark_aliases and room_number:
        landmark_aliases = [room_number, f"{room_number}房间", f"Room {room_number}"]
    landmark_weight = _to_float(row.get("visual_landmark_weight"), default=4.0) or 4.0
    return (
        [{"aliases": landmark_aliases, "weight": landmark_weight}]
        if landmark_aliases
        else []
    )


def build_node(row: dict[str, str], row_number: int) -> dict[str, Any] | None:
    node_id = (row.get("id") or "").strip()
    if not node_id:
        return None

    room_number = (row.get("room_number") or "").strip()
    name_zh = (row.get("name_zh") or room_number or node_id).strip()
    name_en = (row.get("name_en") or room_number or node_id).strip()
    floor = _to_int(row.get("floor"), default=4)
    kind = (row.get("kind") or "room").strip()
    if floor is None:
        raise ValueError(f"第 {row_number} 行缺少 floor")

    aliases = _split_items(row.get("aliases"))
    for candidate in (room_number, name_zh, name_en):
        if candidate and candidate not in aliases and candidate != node_id:
            aliases.append(candidate)

    node: dict[str, Any] = {
        "id": node_id,
        "name_zh": name_zh,
        "name_en": name_en,
        "floor": floor,
        "kind": kind,
        "aliases": aliases,
        "visual_landmarks": _landmarks_from_row(row, room_number),
        "remark": (row.get("remark") or "").strip(),
    }

    if room_number:
        node["room_number"] = room_number

    x_m = _to_float(row.get("x_m"))
    y_m = _to_float(row.get("y_m"))
    x_ratio = _to_float(row.get("x_ratio"))
    y_ratio = _to_float(row.get("y_ratio"))
    if any(value is not None for value in (x_m, y_m, x_ratio, y_ratio)):
        node["position"] = {
            "x_m": x_m,
            "y_m": y_m,
            "x_ratio": x_ratio,
            "y_ratio": y_ratio,
            "estimated": True,
        }

    source = (row.get("source") or "").strip()
    if source:
        node["source"] = source

    return node


def build_edge(
    row: dict[str, str], row_number: int, node_ids: set[str]
) -> dict[str, Any] | None:
    start = (row.get("from") or "").strip()
    end = (row.get("to") or "").strip()
    if not start and not end:
        return None
    if start not in node_ids or end not in node_ids:
        raise ValueError(f"第 {row_number} 行边引用未知节点：{start} -> {end}")

    distance = _to_float(row.get("distance_m"))
    if distance is None or distance <= 0:
        raise ValueError(f"第 {row_number} 行边距离必须大于 0：{start} -> {end}")

    edge: dict[str, Any] = {
        "from": start,
        "to": end,
        "distance_m": distance,
        "bidirectional": _to_bool(row.get("bidirectional"), default=True),
        "accessible": _to_bool(row.get("accessible"), default=True),
    }

    instructions = _json_value(row.get("instructions_json"), default=None)
    if instructions is None:
        instructions = {}
        forward = {
            key: (row.get(f"forward_{key}") or "").strip()
            for key in ("zh", "en")
        }
        reverse = {
            key: (row.get(f"reverse_{key}") or "").strip()
            for key in ("zh", "en")
        }
        forward = {key: value for key, value in forward.items() if value}
        reverse = {key: value for key, value in reverse.items() if value}
        if forward:
            instructions["forward"] = forward
        if reverse:
            instructions["reverse"] = reverse
    if instructions:
        if not isinstance(instructions, dict):
            raise ValueError(f"第 {row_number} 行 instructions_json 必须是对象")
        edge["instructions"] = instructions

    return edge


def build_metadata(
    metadata_rows: list[dict[str, str]], existing: dict[str, Any]
) -> dict[str, Any]:
    metadata: dict[str, Any] = {}
    for row in metadata_rows:
        key = (row.get("key") or "").strip()
        if not key:
            continue
        metadata[key] = _json_value(row.get("value_json"), default=row.get("value_json"))

    return {
        "schema_version": metadata.get("schema_version", existing.get("schema_version", 1)),
        "building": metadata.get(
            "building",
            existing.get(
                "building",
                {
                    "id": "LONGBIN",
                    "name_zh": "龙宾楼",
                    "name_en": "Longbin Building",
                },
            ),
        ),
        "floor_plan": metadata.get(
            "floor_plan", existing.get("floor_plan", DEFAULT_FLOOR_PLAN)
        ),
    }


def build_map(
    node_rows: list[dict[str, str]],
    edge_rows: list[dict[str, str]],
    metadata_rows: list[dict[str, str]],
    output: Path,
) -> dict[str, Any]:
    existing: dict[str, Any] = {}
    if output.exists():
        with output.open(encoding="utf-8") as file:
            loaded = json.load(file)
            if isinstance(loaded, dict):
                existing = loaded

    nodes: list[dict[str, Any]] = []
    seen: set[str] = set()
    for index, row in enumerate(node_rows, start=2):
        node = build_node(row, index)
        if node is None:
            continue
        node_id = node["id"]
        if node_id in seen:
            raise ValueError(f"节点 ID 重复：{node_id}")
        seen.add(node_id)
        nodes.append(node)

    if not nodes:
        raise ValueError("没有可导入的节点")

    edges: list[dict[str, Any]] = []
    node_ids = {node["id"] for node in nodes}
    for index, row in enumerate(edge_rows, start=2):
        edge = build_edge(row, index, node_ids)
        if edge is not None:
            edges.append(edge)

    metadata = build_metadata(metadata_rows, existing)
    return {
        "schema_version": metadata["schema_version"],
        "building": metadata["building"],
        "floor_plan": metadata["floor_plan"],
        "nodes": nodes,
        "edges": edges,
    }


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="将 Excel/CSV 地图表导入 data/building_map.json。"
    )
    parser.add_argument("--input", type=Path, default=DEFAULT_INPUT)
    parser.add_argument("--output", type=Path, default=DEFAULT_OUTPUT)
    parser.add_argument("--sheet", default="nodes", help="兼容旧参数：节点表名称。")
    parser.add_argument("--nodes-sheet", default=None)
    parser.add_argument("--edges-sheet", default="edges")
    parser.add_argument("--metadata-sheet", default="metadata")
    parser.add_argument(
        "--dry-run",
        action="store_true",
        help="只解析并打印节点和边数量，不写入 JSON。",
    )
    return parser.parse_args()


def main() -> int:
    args = parse_args()
    nodes_sheet = args.nodes_sheet or args.sheet
    node_rows = read_rows(args.input, nodes_sheet)
    if args.input.suffix.lower() == ".csv":
        edge_rows: list[dict[str, str]] = []
        metadata_rows: list[dict[str, str]] = []
    else:
        edge_rows = read_rows(args.input, args.edges_sheet)
        metadata_rows = read_rows(args.input, args.metadata_sheet)
    data = build_map(node_rows, edge_rows, metadata_rows, args.output)
    if args.dry_run:
        print(f"parsed_nodes={len(data['nodes'])} parsed_edges={len(data['edges'])}")
        return 0

    args.output.parent.mkdir(parents=True, exist_ok=True)
    with args.output.open("w", encoding="utf-8") as file:
        json.dump(data, file, ensure_ascii=False, indent=2)
        file.write("\n")
    print(
        f"wrote {len(data['nodes'])} nodes and {len(data['edges'])} edges to {args.output}"
    )
    return 0


if __name__ == "__main__":
    try:
        raise SystemExit(main())
    except Exception as exc:
        print(f"error: {exc}", file=sys.stderr)
        raise SystemExit(1)
